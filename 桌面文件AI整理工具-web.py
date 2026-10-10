# -*- coding: utf-8 -*-
"""
文件AI整理工具 - pywebview版（Demo D 温暖柔和风）
后端API层：所有业务逻辑保留，前端用HTML/CSS渲染
"""

import os
import json
import shutil
import base64
import threading
import io
import webview
from datetime import datetime

# ============ 默认配置 ============
# 默认整理目录：当前用户的桌面（可在界面里随时切换）
DEFAULT_FOLDER = os.path.join(os.path.expanduser("~"), "Desktop")
DEFAULT_PROVIDER = "deepseek"

# 服务商预设（URL 与模型名以各家官方文档为准）
PROVIDERS = {
    "deepseek": {
        "name": "DeepSeek（推荐）",
        "base_url": "https://api.deepseek.com",
        "model": "deepseek-chat",
        "needs_key": True, "json_mode": True, "local": False,
        "note": "中文效果好、价格低，国内可直连",
    },
    "moonshot": {
        "name": "Kimi 月之暗面",
        "base_url": "https://api.moonshot.cn/v1",
        "model": "moonshot-v1-8k",
        "needs_key": True, "json_mode": True, "local": False,
        "note": "长文本能力强，国内可直连",
    },
    "zhipu": {
        "name": "智谱 GLM",
        "base_url": "https://open.bigmodel.cn/api/paas/v4",
        "model": "glm-4-flash",
        "needs_key": True, "json_mode": True, "local": False,
        "note": "国内可直连，有免费额度",
    },
    "dashscope": {
        "name": "通义千问",
        "base_url": "https://dashscope.aliyuncs.com/compatible-mode/v1",
        "model": "qwen-plus",
        "needs_key": True, "json_mode": True, "local": False,
        "note": "阿里云，国内可直连",
    },
    "siliconflow": {
        "name": "硅基流动",
        "base_url": "https://api.siliconflow.cn/v1",
        "model": "Qwen/Qwen2.5-7B-Instruct",
        "needs_key": True, "json_mode": True, "local": False,
        "note": "一个 Key 可用多种开源模型",
    },
    "openrouter": {
        "name": "OpenRouter",
        "base_url": "https://openrouter.ai/api/v1",
        "model": "deepseek/deepseek-chat",
        "needs_key": True, "json_mode": True, "local": False,
        "note": "一个 Key 可用多家模型（含 GPT / Claude）",
    },
    "openai": {
        "name": "OpenAI",
        "base_url": "https://api.openai.com/v1",
        "model": "gpt-4o-mini",
        "needs_key": True, "json_mode": True, "local": False,
        "note": "需要能访问外网",
    },
    "ollama": {
        "name": "本地模型（隐私模式）",
        "base_url": "http://localhost:11434/v1",
        "model": "qwen2.5:7b",
        "needs_key": False, "json_mode": False, "local": True,
        "note": "AI 完全在本机运行，文件内容不会上传；需先安装并启动 Ollama",
    },
    "custom": {
        "name": "自定义（OpenAI 兼容）",
        "base_url": "",
        "model": "",
        "needs_key": True, "json_mode": False, "local": False,
        "note": "任何 OpenAI 兼容接口：填入接口地址和模型名即可",
    },
}

# ============ 本地配置（API Key、历史分类等，存于用户目录，不入仓库） ============
APP_DIR = os.path.join(os.path.expanduser("~"), ".fileai")
CONFIG_FILE = os.path.join(APP_DIR, "config.json")
HISTORY_FILE = os.path.join(APP_DIR, "categories_history.json")
UNDO_FILE = os.path.join(APP_DIR, "last_archive.json")
INDEX_HTML = os.path.join(os.path.dirname(os.path.abspath(__file__)), "index.html")


def _blank_provider_settings(name):
    spec = PROVIDERS.get(name, {})
    return {"base_url": spec.get("base_url", ""), "model": spec.get("model", ""), "api_key": ""}


def normalize_config(raw):
    """把历史/缺省配置统一成 {current, providers} 结构；老配置自动迁移"""
    if not isinstance(raw, dict):
        raw = {}
    providers = raw.get("providers")
    if isinstance(providers, dict) and providers:
        for name in PROVIDERS:
            if not isinstance(providers.get(name), dict):
                providers[name] = _blank_provider_settings(name)
        current = raw.get("current")
        if current not in PROVIDERS:
            current = DEFAULT_PROVIDER
        return {"current": current, "providers": providers}
    cfg = {"current": DEFAULT_PROVIDER,
           "providers": {name: _blank_provider_settings(name) for name in PROVIDERS}}
    old_key = raw.get("api_key")
    if isinstance(old_key, str) and old_key.strip():
        cfg["providers"][DEFAULT_PROVIDER]["api_key"] = old_key.strip()
    return cfg


def load_config():
    raw = {}
    try:
        if os.path.exists(CONFIG_FILE):
            with open(CONFIG_FILE, "r", encoding="utf-8") as f:
                raw = json.load(f)
    except Exception:
        raw = {}
    return normalize_config(raw)


def save_config(cfg):
    try:
        os.makedirs(APP_DIR, exist_ok=True)
        with open(CONFIG_FILE, "w", encoding="utf-8") as f:
            json.dump(cfg, f, ensure_ascii=False, indent=2)
        return True
    except Exception:
        return False


def get_active_settings(cfg):
    """取出当前生效的服务商完整配置"""
    name = cfg.get("current", DEFAULT_PROVIDER)
    spec = PROVIDERS.get(name, {})
    entry = (cfg.get("providers") or {}).get(name, {})
    return {
        "provider": name,
        "name": spec.get("name", name),
        "base_url": (entry.get("base_url") or spec.get("base_url") or "").strip(),
        "model": (entry.get("model") or spec.get("model") or "").strip(),
        "api_key": (entry.get("api_key") or "").strip(),
        "needs_key": spec.get("needs_key", True),
        "json_mode": spec.get("json_mode", False),
        "local": spec.get("local", False),
    }


def settings_problem(settings):
    """配置是否可用；可用返回 None，否则返回给用户看的提示"""
    spec = PROVIDERS.get(settings.get("provider"), {})
    if not settings.get("base_url"):
        return "还没填接口地址，请到「API 设置」里补上"
    if not settings.get("model"):
        return "还没填模型名，请到「API 设置」里补上"
    if spec.get("needs_key", True) and not settings.get("api_key"):
        return "还没填 API Key，请到「API 设置」里补上"
    return None

SKIP_EXTENSIONS = {".lnk", ".exe", ".nir", ".dll", ".sys", ".tmp", ".crdownload"}
TEXT_EXTENSIONS = {".txt", ".csv", ".r", ".py", ".md", ".json", ".xml", ".html", ".htm"}
IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".gif", ".bmp", ".webp", ".tiff"}
MAX_CONTENT_CHARS = 3000

# 文件名/文件夹名清理
INVALID_NAME_CHARS = '<>:"/\\|?*'
WINDOWS_RESERVED_NAMES = (
    {"CON", "PRN", "AUX", "NUL"}
    | {f"COM{i}" for i in range(1, 10)}
    | {f"LPT{i}" for i in range(1, 10)}
)
MAX_NAME_LEN = 120
SEARCH_MAX_CANDIDATES = 60


def sanitize_filename(name):
    """清理文件名/文件夹名：去掉非法与危险字符、限长、规避 Windows 保留名。
    返回清理后的名字（不含扩展名）；清理后为空则返回空字符串。"""
    name = str(name or "")
    name = "".join(ch for ch in name if ord(ch) >= 32)
    for ch in INVALID_NAME_CHARS:
        name = name.replace(ch, "_")
    name = name.strip().strip(".").strip()
    if len(name) > MAX_NAME_LEN:
        name = name[:MAX_NAME_LEN].rstrip()
    if not name:
        return ""
    if name.upper() in WINDOWS_RESERVED_NAMES:
        name = "_" + name
    return name


# ============ 文件内容读取（复用原逻辑） ============
def read_file_content(filepath):
    ext = os.path.splitext(filepath)[1].lower()
    try:
        if ext in IMAGE_EXTENSIONS:
            return "[图片文件，请点右上角'打开原文件'用看图软件查看]"
        if ext == ".pdf":
            import pymupdf
            doc = pymupdf.open(filepath)
            try:
                text = ""
                for i in range(min(3, doc.page_count)):
                    text += doc[i].get_text()
            finally:
                doc.close()
            return text[:MAX_CONTENT_CHARS] if text.strip() else "[PDF无文字内容（可能是扫描件）]"
        if ext == ".docx":
            from docx import Document
            doc = Document(filepath)
            text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
            return text[:MAX_CONTENT_CHARS] if text.strip() else "[空文档或无文字内容]"
        elif ext == ".pptx":
            from pptx import Presentation
            prs = Presentation(filepath)
            text = ""
            for idx, slide in enumerate(prs.slides):
                slide_text = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_text.append(shape.text.strip())
                if slide_text:
                    text += f"--- 幻灯片 {idx+1} ---\n" + "\n".join(slide_text) + "\n\n"
            return text[:MAX_CONTENT_CHARS] if text.strip() else "[PPT无文字内容]"
        elif ext == ".xlsx":
            from openpyxl import load_workbook
            wb = load_workbook(filepath, read_only=True, data_only=True)
            text = ""
            for sheet_name in wb.sheetnames[:3]:
                ws = wb[sheet_name]
                text += f"=== Sheet: {sheet_name} ===\n"
                for i, row in enumerate(ws.iter_rows(values_only=True)):
                    if i >= 15:
                        break
                    row_text = [str(c) if c is not None else "" for c in row]
                    if any(row_text):
                        text += " | ".join(row_text) + "\n"
                text += "\n"
            wb.close()
            return text[:MAX_CONTENT_CHARS] if text.strip() else "[Excel无数据]"
        elif ext in TEXT_EXTENSIONS:
            for enc in ["utf-8", "gbk", "gb2312", "latin-1"]:
                try:
                    with open(filepath, "r", encoding=enc) as f:
                        return f.read(MAX_CONTENT_CHARS)
                except (UnicodeDecodeError, UnicodeError):
                    continue
            return ""
        elif ext == ".doc":
            return "[旧版Word文档(.doc)，无法自动读取，请点右上角'打开原文件'用WPS查看]"
        else:
            return f"[文件类型 {ext}，暂不支持预览]"
    except Exception as e:
        return f"[读取失败: {str(e)[:100]}]"


def render_pdf_first_page_b64(filepath, max_width=600):
    """渲染PDF第一页为base64 data URL"""
    try:
        import pymupdf
        doc = pymupdf.open(filepath)
        page = doc[0]
        zoom = max_width / page.rect.width
        mat = pymupdf.Matrix(zoom, zoom)
        pix = page.get_pixmap(matrix=mat)
        img_bytes = pix.tobytes("png")
        doc.close()
        b64 = base64.b64encode(img_bytes).decode("ascii")
        return f"data:image/png;base64,{b64}"
    except Exception:
        return None


# ============ AI 请求（统一入口，兼容各家 OpenAI 兼容接口） ============
class AIError(Exception):
    """带用户可读提示的 AI 调用错误"""


def chat_endpoint(base_url):
    return base_url.rstrip("/") + "/chat/completions"


def describe_http_error(resp):
    code = resp.status_code
    detail = ""
    try:
        body = resp.json()
        err = body.get("error")
        if isinstance(err, dict):
            detail = str(err.get("message") or "")
        elif isinstance(err, str):
            detail = err
        if not detail:
            detail = str(body.get("message") or "")
    except Exception:
        detail = (resp.text or "")[:120]
    detail = detail.strip()[:120]
    if code in (401, 403):
        base = "API Key 无效或没有权限"
    elif code == 404:
        base = "接口地址或模型名不对"
    elif code == 429:
        base = "请求太频繁（被限流），请稍后再试"
    elif 500 <= code < 600:
        base = "服务商暂时故障，请稍后重试"
    else:
        base = f"服务商返回错误 {code}"
    return f"{base}（{detail}）" if detail else base


def parse_json_reply(text):
    """尽力从模型回复里解析 JSON（兼容本地小模型不规范的输出）"""
    if not text:
        return None
    text = text.strip()
    try:
        return json.loads(text)
    except Exception:
        pass
    if "```" in text:
        for seg in text.split("```"):
            seg = seg.strip()
            if seg[:4].lower() == "json":
                seg = seg[4:].strip()
            if seg.startswith("{"):
                try:
                    return json.loads(seg)
                except Exception:
                    pass
    i, j = text.find("{"), text.rfind("}")
    if i != -1 and j > i:
        try:
            return json.loads(text[i:j + 1])
        except Exception:
            pass
    return None


def http_post(url, headers, payload, timeout, local=False):
    """local=True 时不走系统代理——本地模型常被系统代理拦掉"""
    import requests
    if local:
        session = requests.Session()
        session.trust_env = False
        return session.post(url, headers=headers, json=payload, timeout=timeout)
    return requests.post(url, headers=headers, json=payload, timeout=timeout)


def http_get(url, timeout, local=False):
    import requests
    if local:
        session = requests.Session()
        session.trust_env = False
        return session.get(url, timeout=timeout)
    return requests.get(url, timeout=timeout)


def chat_json(settings, prompt, timeout=90):
    """统一调用 OpenAI 兼容的 /chat/completions，返回解析后的 dict"""
    problem = settings_problem(settings)
    if problem:
        raise AIError(problem)
    import requests
    headers = {"Content-Type": "application/json"}
    if settings.get("api_key"):
        headers["Authorization"] = f"Bearer {settings['api_key']}"
    payload = {
        "model": settings["model"],
        "messages": [{"role": "user", "content": prompt}],
        "temperature": 0.3,
    }
    if settings.get("json_mode"):
        payload["response_format"] = {"type": "json_object"}
    try:
        resp = http_post(chat_endpoint(settings["base_url"]), headers, payload, timeout, local=settings.get("local"))
    except requests.exceptions.ConnectionError:
        if settings.get("local"):
            raise AIError("连不上本地模型服务，请确认 Ollama 已启动")
        raise AIError("网络连接失败，请检查网络或接口地址")
    except requests.exceptions.Timeout:
        raise AIError("请求超时（模型太慢或网络不通）")
    except Exception as e:
        raise AIError(f"请求失败：{str(e)[:80]}")
    if resp.status_code != 200:
        raise AIError(describe_http_error(resp))
    try:
        content = resp.json()["choices"][0]["message"]["content"]
    except Exception:
        raise AIError("服务商返回的数据格式不对，请检查接口地址")
    parsed = parse_json_reply(content)
    if parsed is None:
        raise AIError("模型没有返回合法 JSON（本地小模型常见，建议换更大的模型）")
    return parsed


# ============ AI分析（复用原逻辑） ============
def analyze_file_with_ai(filename, content, folder_list, settings):
    folder_str = "\n".join([f"- {f}" for f in folder_list])
    prompt = f"""你是一个文件整理助手。请根据以下信息，帮我给这个文件起一个清晰的名字，并判断它应该归到哪个分类文件夹。

【文件名】{filename}

【文件内容摘要】
{content[:2000]}

【可选的分类文件夹列表】
{folder_str}

请严格按照以下JSON格式返回（不要输出其他内容）：
{{
  "new_filename": "建议的新文件名，不含扩展名，格式建议：项目/主题-内容描述-版本/日期",
  "target_folder": "从上面的分类列表中选一个最合适的，必须完全匹配列表中的名称",
  "reason": "一句话说明为什么这么归类"
}}

注意：
1. new_filename不要包含扩展名，不要包含特殊字符 \\/ : * ? " < > |
2. target_folder：优先选最匹配的现有分类；但**不要硬塞到不太相关的文件夹里**。如果文件的主题和上面列出的分类都不太匹配，就根据文件内容自己起一个简洁准确的新分类名（比如"引物设计"、"海南数据"、"文献阅读"），这是完全正常且鼓励的
3. 新建分类名要概括文件的核心主题，不要太泛（比如不要只写"资料"），也不要太碎
4. 如果文件内容无法判断，根据文件名做合理推测
"""
    parsed = chat_json(settings, prompt, timeout=90)
    return parsed.get("new_filename", filename), parsed.get("target_folder", ""), parsed.get("reason", "")


def scan_folder(root_path):
    files = []
    folders = []
    if not os.path.exists(root_path):
        return files, folders
    all_items = os.listdir(root_path)
    folders = [item for item in all_items if os.path.isdir(os.path.join(root_path, item))]
    for item in all_items:
        filepath = os.path.join(root_path, item)
        if os.path.isfile(filepath):
            ext = os.path.splitext(item)[1].lower()
            if ext not in SKIP_EXTENSIONS and not item.startswith("~$"):
                files.append(item)
    return files, folders


def scan_all_files_recursive(root_path):
    results = []
    for dirpath, dirnames, filenames in os.walk(root_path):
        for f in filenames:
            ext = os.path.splitext(f)[1].lower()
            if ext not in SKIP_EXTENSIONS and not f.startswith("~$"):
                full_path = os.path.join(dirpath, f)
                rel_path = os.path.relpath(full_path, root_path)
                results.append((rel_path, full_path))
    return results


def execute_archive(results, root_path, keep_original_name=False):
    log = []
    moves = []
    for item in results:
        if not item["selected"]:
            continue
        old_path = os.path.join(root_path, item["original_name"])
        ext = os.path.splitext(item["original_name"])[1]
        if keep_original_name:
            new_name = item["original_name"]
        else:
            suggested = sanitize_filename(item.get("new_name"))
            # 建议名为空（或清理后为空）时保留原文件名（清空建议名 = 不改名）
            new_name = (suggested + ext) if suggested else item["original_name"]
        target_folder = sanitize_filename(item.get("target_folder"))
        if not target_folder:
            log.append(f"[跳过] {item['original_name']} - 未指定目标文件夹")
            continue
        target_dir = os.path.join(root_path, target_folder)
        if not os.path.exists(target_dir):
            try:
                os.makedirs(target_dir, exist_ok=True)
                log.append(f"[新建文件夹] {target_folder}")
            except Exception as e:
                log.append(f"[失败] {item['original_name']} - 创建文件夹失败: {str(e)[:60]}")
                continue
        new_path = os.path.join(target_dir, new_name)
        base_stem = os.path.splitext(new_name)[0]
        counter = 1
        while os.path.exists(new_path):
            new_name = f"{base_stem}_{counter}{ext}"
            new_path = os.path.join(target_dir, new_name)
            counter += 1
        try:
            shutil.move(old_path, new_path)
            moves.append({"from": old_path, "to": new_path})
            log.append(f"[成功] {item['original_name']} → {target_folder}/{new_name}")
        except Exception as e:
            log.append(f"[失败] {item['original_name']} - {str(e)[:80]}")
    return log, moves


# ============ API类（暴露给前端） ============
class Api:
    def __init__(self):
        self.window = None
        self.current_folder = DEFAULT_FOLDER
        self.files = []
        self.folders = []
        self.results = []
        self.user_categories = []
        self.ai_categories = []
        self.is_analyzing = False
        self.is_archiving = False
        self._config = load_config()
        self._load_history()

    def _load_history(self):
        """读取按文件夹分组的 AI 分类历史（不同文件夹互相独立）"""
        self._history = {}
        try:
            if os.path.exists(HISTORY_FILE):
                with open(HISTORY_FILE, "r", encoding="utf-8") as f:
                    data = json.load(f)
                if isinstance(data, dict):
                    self._history = data
                elif isinstance(data, list):
                    # 兼容旧的全局列表格式：归到当前文件夹
                    self._history = {self.current_folder: data} if data else {}
        except Exception:
            self._history = {}
        self.ai_categories = list(self._history.get(self.current_folder, []))

    def _save_history(self):
        """保存当前文件夹的分类到历史文件"""
        try:
            self._history[self.current_folder] = list(self.ai_categories)
            os.makedirs(APP_DIR, exist_ok=True)
            with open(HISTORY_FILE, "w", encoding="utf-8") as f:
                json.dump(self._history, f, ensure_ascii=False, indent=2)
        except Exception:
            pass

    def _push(self, js_call):
        """安全地向前端推送JS调用"""
        if self.window:
            try:
                self.window.evaluate_js(js_call)
            except Exception:
                pass

    def _log(self, msg):
        ts = datetime.now().strftime("%H:%M:%S")
        self._push(f"window.on_log({json.dumps(f'[{ts}] {msg}', ensure_ascii=False)})")

    def _status(self, text):
        self._push(f"window.on_status({json.dumps(text, ensure_ascii=False)})")

    def _settings(self):
        """当前生效的服务商配置"""
        return get_active_settings(self._config)

    def _is_configured(self):
        return settings_problem(self._settings()) is None

    # ---- 前端调用的方法 ----
    def init(self):
        """初始化，返回当前状态"""
        return {
            "current_folder": self.current_folder,
            "ai_categories": self.ai_categories,
            "api_configured": self._is_configured(),
        }

    def choose_folder(self):
        """弹原生对话框选文件夹"""
        result = self.window.create_file_dialog(webview.FOLDER_DIALOG)
        if result and len(result) > 0:
            new_folder = result[0]
            self._save_history()  # 先保存当前文件夹的分类
            self.current_folder = new_folder
            self.user_categories = []
            self.ai_categories = list(self._history.get(new_folder, []))
            self.results = []
            self._log(f"切换整理文件夹到: {new_folder}")
            self._log(f"该文件夹已有的AI历史分类：{len(self.ai_categories)}个")
            return {"ok": True, "folder": new_folder}
        return {"ok": False}

    def scan(self):
        """扫描当前文件夹的散落文件"""
        self._status("正在扫描...")
        self.files, self.folders = scan_folder(self.current_folder)
        self.files.sort(key=lambda f: (os.path.splitext(f)[1].lower(), f))
        self.results = []
        file_list = []
        for f in self.files:
            item = {
                "original_name": f,
                "new_name": os.path.splitext(f)[0],
                "target_folder": "",
                "reason": "待分析",
                "selected": True,
            }
            self.results.append(item)
            file_list.append(item)
        self._log(f"扫描完成：发现 {len(self.files)} 个散落文件，{len(self.folders)} 个子文件夹（已按文件格式排序）")
        self._status(f"扫描完成，共 {len(self.files)} 个文件")
        return {"files": file_list, "folders": self.folders}

    def analyze(self):
        """后台跑AI分析"""
        problem = settings_problem(self._settings())
        if problem:
            return {"ok": False, "msg": problem}
        if self.is_analyzing or not self.results:
            return {"ok": False, "msg": "正在分析或没有文件"}
        self.is_analyzing = True
        threading.Thread(target=self._analyze_worker, daemon=True).start()
        return {"ok": True}

    def _analyze_worker(self):
        total = len(self.results)
        folder_set = set(self.folders)
        folder_set.update(self.user_categories)
        folder_set.update(self.ai_categories)
        folder_list = sorted(folder_set)
        self._log(f"分类候选：共{len(folder_list)}个（现有{len(self.folders)} + 自定义{len(self.user_categories)} + AI积累{len(self.ai_categories)}）")

        for i, item in enumerate(self.results):
            try:
                self._status(f"AI分析中... ({i+1}/{total}) {item['original_name']}")
                filepath = os.path.join(self.current_folder, item["original_name"])
                content = read_file_content(filepath)
                new_name, target_folder, reason = analyze_file_with_ai(item["original_name"], content, folder_list, self._settings())
                item["new_name"] = new_name
                item["target_folder"] = target_folder
                item["reason"] = reason
                self._push(f"window.on_analyze_row({i}, {json.dumps(item, ensure_ascii=False)})")
                self._log(f"分析完成({i+1}/{total}): {item['original_name']} → {item['target_folder']}/{item['new_name']}")
            except Exception as e:
                item["reason"] = f"分析失败: {str(e)[:60]}"
                self._push(f"window.on_analyze_row({i}, {json.dumps(item, ensure_ascii=False)})")
                self._log(f"分析失败({i+1}/{total}): {item['original_name']} - {str(e)[:60]}")
            self._push(f"window.on_progress({i+1}, {total})")

        # 分析完成：按归属文件夹排序
        def sort_key(it):
            f = it["target_folder"]
            if not f:
                return (1, "ZZZZZ")
            return (0, f)
        self.results.sort(key=sort_key)

        # 收集AI新建分类
        existing = set(self.folders) | set(self.user_categories)
        new_cats = set()
        for r in self.results:
            tf = r.get("target_folder", "")
            if tf and tf not in existing and tf not in self.ai_categories:
                new_cats.add(tf)
        if new_cats:
            self.ai_categories.extend(sorted(new_cats))
            self._save_history()
            self._log(f"AI新建了{len(new_cats)}个分类：{', '.join(sorted(new_cats))}（已保存）")

        self.is_analyzing = False
        self._push(f"window.on_analyze_done({json.dumps(self.results, ensure_ascii=False)})")
        self._status(f"分析完成，共 {len(self.results)} 个文件，已按归属文件夹排序")
        self._log("全部分析完成！")

    def get_preview(self, filename):
        """获取文件预览：PDF返回图片dataURL，其他返回文字"""
        filepath = os.path.join(self.current_folder, filename)
        if not os.path.exists(filepath):
            return {"type": "text", "content": "[文件不存在]"}
        ext = os.path.splitext(filename)[1].lower()
        if ext == ".pdf":
            img = render_pdf_first_page_b64(filepath, max_width=600)
            if img:
                return {"type": "image", "content": img}
            return {"type": "text", "content": "[PDF预览失败，请点右上角打开原文件]"}
        else:
            content = read_file_content(filepath)
            return {"type": "text", "content": content}

    def open_file(self, filename):
        filepath = os.path.join(self.current_folder, filename)
        if os.path.exists(filepath):
            try:
                os.startfile(filepath)
                self._log(f"已用默认程序打开: {filename}")
                return {"ok": True}
            except Exception as e:
                return {"ok": False, "msg": str(e)}
        return {"ok": False, "msg": "文件不存在"}

    def open_path(self, full_path):
        """打开搜索结果里的完整路径文件"""
        if os.path.exists(full_path):
            try:
                os.startfile(full_path)
                return {"ok": True}
            except Exception as e:
                return {"ok": False, "msg": str(e)}
        return {"ok": False}

    def search(self, query):
        """后台搜索，通过回调返回结果"""
        problem = settings_problem(self._settings())
        if problem:
            return {"ok": False, "msg": problem}
        if not query or not query.strip():
            return {"ok": False, "msg": "请输入搜索内容"}
        self._log(f"开始搜索：{query}")
        threading.Thread(target=self._search_worker, args=(query.strip(),), daemon=True).start()
        return {"ok": True}

    def _pick_search_candidates(self, query, all_files):
        """本地预筛候选：字面命中的优先，命中不足则退回最近修改的文件。
        这样既能保证关键词命中的文件一定进入候选，又不会把整份列表塞给 AI。"""
        keywords = [k for k in query.replace("，", " ").replace(",", " ").split() if k]
        scored = []
        for rel, full in all_files:
            base = os.path.basename(rel).lower()
            rel_low = rel.lower()
            score = 0
            for kw in keywords:
                kw_low = kw.lower()
                if kw_low in base:
                    score += 3
                elif kw_low in rel_low:
                    score += 1
            if score > 0:
                scored.append((score, rel, full))
        if scored:
            scored.sort(key=lambda x: (-x[0], x[1]))
            return [(rel, full) for _, rel, full in scored[:SEARCH_MAX_CANDIDATES]]

        def _mtime(pair):
            try:
                return os.path.getmtime(pair[1])
            except OSError:
                return 0
        return sorted(all_files, key=_mtime, reverse=True)[:SEARCH_MAX_CANDIDATES]

    def _search_worker(self, query):
        try:
            all_files = scan_all_files_recursive(self.current_folder)
            if not all_files:
                self._push(f"window.on_search_results({json.dumps(query, ensure_ascii=False)}, [])")
                return
            candidates = self._pick_search_candidates(query, all_files)
            file_list_str = "\n".join([f"{i+1}. {rel}" for i, (rel, _) in enumerate(candidates)])
            prompt = f"""用户要找一个文件，搜索意图是："{query}"

以下是候选文件列表（相对路径，共{len(candidates)}个）：
{file_list_str}

请从上面的候选列表中，找出所有符合用户搜索意图的文件，按匹配度从高到低排序，最多返回10个。
严格返回JSON格式：
{{
  "results": [
    {{"index": 1, "reason": "为什么匹配"}},
    ...
  ]
}}
index是文件在上面列表里的序号（从1开始）。匹配度高的排前面，明显不相关的不要列进来；如果都不相关，results 返回空数组。
"""
            parsed = chat_json(self._settings(), prompt, timeout=60)
            matched = []
            for item in parsed.get("results", []):
                idx = item.get("index", 0) - 1
                if 0 <= idx < len(candidates):
                    rel_path, full_path = candidates[idx]
                    matched.append({"path": rel_path, "full_path": full_path, "reason": item.get("reason", "")})
            self._log(f"搜索完成：'{query}'，候选{len(candidates)}个，命中{len(matched)}个")
            self._push(f"window.on_search_results({json.dumps(query, ensure_ascii=False)}, {json.dumps(matched, ensure_ascii=False)})")
        except Exception as e:
            self._log(f"搜索失败: {str(e)[:100]}")
            self._push(f"window.on_search_results({json.dumps(query, ensure_ascii=False)}, [])")

    def archive(self, keep_name=False):
        """执行归档"""
        if self.is_archiving:
            return {"ok": False}
        selected_count = sum(1 for r in self.results if r["selected"])
        if selected_count == 0:
            return {"ok": False, "msg": "没有选中任何文件"}
        self.is_archiving = True
        self._status("正在归档...")
        threading.Thread(target=self._archive_worker, args=(keep_name,), daemon=True).start()
        return {"ok": True, "count": selected_count}

    def _archive_worker(self, keep_name):
        log, moves = execute_archive(self.results, self.current_folder, keep_name)
        self.is_archiving = False
        success = sum(1 for l in log if l.startswith("[成功]"))
        for line in log:
            self._log(line)
        self._log("=" * 50)
        self._log(f"归档完成！成功 {success} 个文件。")
        if moves:
            self._save_undo(moves)
            self._log("可点「撤销上次归档」还原本次操作")
        self._push(f"window.on_archive_done({json.dumps(log, ensure_ascii=False)}, {success})")
        self._status(f"归档完成：成功 {success} 个")

    def _save_undo(self, moves):
        """记录本次归档，供撤销使用"""
        try:
            os.makedirs(APP_DIR, exist_ok=True)
            with open(UNDO_FILE, "w", encoding="utf-8") as f:
                json.dump(
                    {"time": datetime.now().strftime("%Y-%m-%d %H:%M:%S"), "moves": moves},
                    f, ensure_ascii=False, indent=2,
                )
        except Exception:
            pass

    def can_undo(self):
        """是否有可撤销的归档记录"""
        return {"can_undo": os.path.exists(UNDO_FILE)}

    def undo_archive(self):
        """撤销上次归档：把文件移回原位"""
        if not os.path.exists(UNDO_FILE):
            return {"ok": False, "msg": "没有可撤销的归档记录"}
        try:
            with open(UNDO_FILE, "r", encoding="utf-8") as f:
                record = json.load(f)
        except Exception as e:
            return {"ok": False, "msg": f"读取撤销记录失败: {str(e)[:60]}"}
        restored, skipped = 0, 0
        for mv in reversed(record.get("moves", [])):
            src = mv.get("to")
            dst = mv.get("from")
            if not src or not dst or not os.path.exists(src) or os.path.exists(dst):
                skipped += 1
                continue
            try:
                parent = os.path.dirname(dst)
                if parent:
                    os.makedirs(parent, exist_ok=True)
                shutil.move(src, dst)
                restored += 1
            except Exception:
                skipped += 1
        try:
            os.remove(UNDO_FILE)
        except Exception:
            pass
        self._log(f"已撤销上次归档：还原 {restored} 个文件" + (f"，跳过 {skipped} 个" if skipped else ""))
        return {"ok": True, "restored": restored, "skipped": skipped}

    def get_categories(self):
        """返回三层分类"""
        return {
            "existing": self.folders,
            "user": self.user_categories,
            "ai": self.ai_categories,
        }

    def set_user_categories(self, lines):
        """保存用户自定义分类"""
        if lines:
            self.user_categories = [l.strip() for l in lines if l.strip()]
        else:
            self.user_categories = []
        if self.user_categories:
            self._log(f"已设置自定义分类体系（{len(self.user_categories)}个分类）")
        else:
            self._log("已清空自定义分类")
        return {"ok": True, "user": self.user_categories}

    def get_provider_config(self):
        """返回服务商列表与当前配置（绝不回传明文 Key）"""
        cfg = self._config
        current = cfg.get("current", DEFAULT_PROVIDER)
        providers = []
        for name, spec in PROVIDERS.items():
            entry = (cfg.get("providers") or {}).get(name, {})
            providers.append({
                "key": name,
                "name": spec.get("name", name),
                "note": spec.get("note", ""),
                "needs_key": spec.get("needs_key", True),
                "local": spec.get("local", False),
                "preset_base_url": spec.get("base_url", ""),
                "preset_model": spec.get("model", ""),
                "saved_base_url": (entry.get("base_url") or "").strip(),
                "saved_model": (entry.get("model") or "").strip(),
                "has_key": bool((entry.get("api_key") or "").strip()),
            })
        active = get_active_settings(cfg)
        return {
            "providers": providers,
            "current": current,
            "base_url": active["base_url"],
            "model": active["model"],
            "has_key": bool(active["api_key"]),
            "configured": settings_problem(active) is None,
        }

    def set_provider_config(self, provider, base_url, model, api_key):
        """保存服务商配置；api_key 留空表示沿用已保存的"""
        if provider not in PROVIDERS:
            return {"ok": False, "msg": "未知的服务商"}
        spec = PROVIDERS[provider]
        entry = self._config.setdefault("providers", {}).setdefault(provider, _blank_provider_settings(provider))
        entry["base_url"] = (base_url or "").strip() or spec.get("base_url", "")
        entry["model"] = (model or "").strip() or spec.get("model", "")
        key = (api_key or "").strip()
        if not spec.get("needs_key", True):
            entry["api_key"] = ""
        elif key:
            entry["api_key"] = key
        self._config["current"] = provider
        if not save_config(self._config):
            return {"ok": False, "msg": "保存失败，请检查目录权限"}
        active = get_active_settings(self._config)
        self._log(f"已切换服务商：{active['name']}")
        return {"ok": True, "configured": settings_problem(active) is None}

    def test_connection(self, provider, base_url, model, api_key):
        """测试当前填写的配置能不能连通（不保存）"""
        if provider not in PROVIDERS:
            return {"ok": False, "msg": "未知的服务商"}
        import requests
        spec = PROVIDERS[provider]
        saved = (self._config.get("providers") or {}).get(provider, {})
        settings = {
            "provider": provider,
            "base_url": (base_url or "").strip() or spec.get("base_url", ""),
            "model": (model or "").strip() or spec.get("model", ""),
            "api_key": ((api_key or "").strip() or (saved.get("api_key") or "")) if spec.get("needs_key", True) else "",
            "needs_key": spec.get("needs_key", True),
            "json_mode": spec.get("json_mode", False),
            "local": spec.get("local", False),
        }
        problem = settings_problem(settings)
        if problem:
            return {"ok": False, "msg": problem}
        headers = {"Content-Type": "application/json"}
        if settings["api_key"]:
            headers["Authorization"] = f"Bearer {settings['api_key']}"
        payload = {"model": settings["model"], "messages": [{"role": "user", "content": "ping"}]}
        try:
            resp = http_post(chat_endpoint(settings["base_url"]), headers, payload, 25, local=settings["local"])
        except requests.exceptions.ConnectionError:
            if settings["local"]:
                return {"ok": False, "msg": "连不上本地模型服务，请确认 Ollama 已启动"}
            return {"ok": False, "msg": "网络连接失败，请检查网络或接口地址"}
        except requests.exceptions.Timeout:
            return {"ok": False, "msg": "连接超时，模型响应太慢或网络不通"}
        except Exception as e:
            return {"ok": False, "msg": f"请求失败：{str(e)[:80]}"}
        if resp.status_code != 200:
            return {"ok": False, "msg": describe_http_error(resp)}
        try:
            resp.json()["choices"][0]["message"]
        except Exception:
            return {"ok": False, "msg": "接口能连通，但返回格式不像 OpenAI 兼容接口"}
        return {"ok": True, "msg": f"连接成功，模型：{settings['model']}"}

    def check_local_model(self, base_url=""):
        """检测本地 Ollama 是否在运行、是否装了目标模型"""
        spec = PROVIDERS["ollama"]
        saved = (self._config.get("providers") or {}).get("ollama", {})
        base = (base_url or "").strip() or (saved.get("base_url") or "").strip() or spec["base_url"]
        root = base.rstrip("/")
        if root.endswith("/v1"):
            root = root[:-3]
        tags_url = root.rstrip("/") + "/api/tags"
        try:
            resp = http_get(tags_url, 3, local=True)
        except Exception:
            return {"running": False, "msg": "没检测到本地模型服务。请先安装 Ollama 并启动它"}
        if resp.status_code != 200:
            return {"running": False, "msg": f"本地服务返回 {resp.status_code}，请确认它是 Ollama"}
        try:
            models = [str(m.get("name", "")) for m in resp.json().get("models", [])]
        except Exception:
            models = []
        want = (saved.get("model") or "").strip() or spec["model"]
        base_name = want.split(":")[0]
        if any(m.split(":")[0] == base_name for m in models):
            return {"running": True, "models": models, "msg": f"已就绪：{want}"}
        return {
            "running": True,
            "models": models,
            "msg": f"服务已启动，但没有模型 {want}。请运行：ollama pull {want}",
        }

    def toggle_select(self, index, selected):
        if 0 <= index < len(self.results):
            self.results[index]["selected"] = selected
            return {"ok": True}
        return {"ok": False}

    def select_all(self):
        for r in self.results:
            r["selected"] = True
        return {"ok": True}

    def deselect_all(self):
        for r in self.results:
            r["selected"] = False
        return {"ok": True}

    def update_result(self, index, field, value):
        """双击编辑：更新new_name或target_folder"""
        if 0 <= index < len(self.results) and field in ("new_name", "target_folder"):
            self.results[index][field] = value
            return {"ok": True}
        return {"ok": False}


def main():
    api = Api()
    window = webview.create_window(
        title="文件AI整理工具",
        url=INDEX_HTML,
        js_api=api,
        width=1320,
        height=820,
        background_color="#F5F1E8",
    )
    api.window = window
    webview.start(debug=False)


if __name__ == "__main__":
    main()
